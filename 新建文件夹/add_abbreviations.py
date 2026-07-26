from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

src = r"F:\Desk\COLREG_v5.docx"
out = r"F:\Desk\COLREG_v5_abbreviations.docx"

entries = [
    ("A*", "A-star search algorithm"),
    ("AI", "Artificial Intelligence"),
    ("AIS", "Automatic Identification System"),
    ("APF", "Artificial Potential Field"),
    ("ASME", "American Society of Mechanical Engineers"),
    ("ASV", "Autonomous Surface Vessel"),
    ("BT", "Behaviour Tree"),
    ("COLREGs", "International Regulations for Preventing Collisions at Sea"),
    ("CP", "Closest Point"),
    ("CPA", "Closest Point of Approach"),
    ("CTRV", "Constant Turn Rate and Velocity"),
    ("DBSCAN", "Density-Based Spatial Clustering of Applications with Noise"),
    ("DCPA", "Distance at Closest Point of Approach"),
    ("EKF", "Extended Kalman Filter"),
    ("GenAI", "Generative Artificial Intelligence"),
    ("GPMP2", "Gaussian Process Motion Planning 2"),
    ("IEEE", "Institute of Electrical and Electronics Engineers"),
    ("IFAC", "International Federation of Automatic Control"),
    ("IMO", "International Maritime Organization"),
    ("IMU", "Inertial Measurement Unit"),
    ("LiDAR", "Light Detection and Ranging"),
    ("LLM", "Large Language Model"),
    ("MASS", "Maritime Autonomous Surface Ship"),
    ("MPC", "Model Predictive Control"),
    ("MSc", "Master of Science"),
    ("PC", "Principal Component"),
    ("PC1", "First Principal Component"),
    ("PC2", "Second Principal Component"),
    ("PCA", "Principal Component Analysis"),
    ("RRT", "Rapidly-exploring Random Tree"),
    ("RPM", "Revolutions Per Minute"),
    ("TCPA", "Time to Closest Point of Approach"),
    ("USV", "Unmanned Surface Vehicle"),
    ("VO", "Velocity Obstacle"),
]

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

d = Document(src)
d.add_page_break()
p = d.add_paragraph()
p.style = d.styles["Heading 1"] if "Heading 1" in d.styles else d.styles["Normal"]
r = p.add_run("Abbreviations and Full Forms")
r.bold = True

table = d.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.autofit = False
widths = [Inches(1.35), Inches(5.9)]
for i, w in enumerate(widths):
    table.columns[i].width = w
hdr = table.rows[0].cells
set_cell_text(hdr[0], "Abbreviation", True)
set_cell_text(hdr[1], "Full form", True)
shade(hdr[0], "D9EAF7")
shade(hdr[1], "D9EAF7")
tr_pr = table.rows[0]._tr.get_or_add_trPr()
repeat = OxmlElement("w:tblHeader")
repeat.set(qn("w:val"), "true")
tr_pr.append(repeat)

for abbr, full in entries:
    cells = table.add_row().cells
    set_cell_text(cells[0], abbr)
    set_cell_text(cells[1], full)
    cells[0].paragraphs[0].alignment = 1

d.save(out)
print(out)
